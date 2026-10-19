"""交付资料与周报/收尾 Word 报告 API（T5-07；FR-P1-08）。"""
from __future__ import annotations

import io
from datetime import datetime, timezone
from uuid import UUID

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from docx import Document
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.db import get_db
from app.models.project import Project
from app.models.stage3 import Issue, Risk, Stakeholder
from app.models.stage5 import Deliverable
from app.models.task import Milestone, Task
from app.schemas.stage5 import DeliverableRead
from app.services import progress as progress_service
from app.services import wbs as wbs_service

router = APIRouter(prefix="/api", tags=["reports"])


def _utcnow() -> datetime:
    return datetime.now(timezone.utc)


def _metrics_of(db: Session, project_id: UUID) -> dict:
    """项目级进度/工时统一口径（与 WBS 页根任务有效进度一致）。"""
    return progress_service.project_metrics(wbs_service.list_tasks_ordered(db, project_id))


def _progress_text(task: Task, effective: dict) -> str:
    """任务行进度文本：父任务取子树汇总后的有效进度，统一 1 位小数。"""
    value = effective.get(task.id, float(task.progress or 0))
    return f"{round(float(value), 1)}%"


@router.get("/projects/{project_id}/reports/weekly-text")
def weekly_text(project_id: UUID, db: Session = Depends(get_db)) -> dict:
    """周报文本（真实数据聚合，带三要素）。"""
    now = _utcnow().astimezone()
    metrics = _metrics_of(db, project_id)
    lines = [
        f"统计截止日期：{now.date().isoformat()}",
        f"数据来源：Project Toolkit（任务/风险/问题/交付清单聚合）",
        f"生成时间：{now.strftime('%Y-%m-%d %H:%M')}",
        "",
        "一、本周进展",
        f"- 任务完成：{metrics['completed_tasks']}/{metrics['total_tasks']}，总工时 {metrics['total_hours']:.1f}h，完成进度 "
        f"{metrics['progress']}%",
        "二、下周计划",
        "- （信息不足：请按实际情况填写）",
        "三、风险与问题",
        f"- 开放风险 {len(list(db.scalars(select(Risk).where(Risk.project_id == project_id, Risk.archived_at.is_(None)))))} 项、"
        f"未关闭问题 {len(list(db.scalars(select(Issue).where(Issue.project_id == project_id, Issue.status != '已关闭'))))} 项",
        "四、需协调事项",
        "- （信息不足：请按实际情况填写）",
    ]
    return {"meta": " | ".join(lines[:3]), "content": "\n".join(lines)}


@router.get("/projects/{project_id}/reports/{report_type}/export")
def export_report(project_id: UUID, report_type: str, db: Session = Depends(get_db)) -> StreamingResponse:
    """Word 导出：weekly=周报 / final=项目收尾报告（python-docx 模板字段映射）。"""
    if report_type not in ("weekly", "final"):
        from app.core.errors import AppError
        raise AppError(422, "invalid_report_type", "仅支持 weekly / final")
    project = db.scalar(select(Project).where(Project.id == project_id))
    if project is None:
        from app.core.errors import AppError
        raise AppError(404, "not_found", "项目不存在")

    if report_type == "weekly":
        document = _build_weekly_doc(db, project)
        filename = f"weekly_report_{project.code}_{_utcnow().date().isoformat()}.docx"
    else:
        document = _build_final_doc(db, project)
        filename = f"final_report_{project.code}_{_utcnow().date().isoformat()}.docx"

    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _build_weekly_doc(db: Session, project: Project) -> Document:
    doc = Document()
    doc.add_heading(f"项目周报：{project.name}", level=1)
    doc.add_paragraph(f"项目编号：{project.code}    统计日期：{_utcnow().astimezone().date().isoformat()}")
    tasks = wbs_service.list_tasks_ordered(db, project.id)
    metrics = progress_service.project_metrics(tasks)
    effective = progress_service.compute_progress(tasks)
    doc.add_heading("一、本周进展", level=2)
    doc.add_paragraph(
        f"任务完成 {metrics['completed_tasks']}/{metrics['total_tasks']}；整体进度 {metrics['progress']}%"
    )
    doc.add_heading("二、任务明细", level=2)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, text in enumerate(["WBS", "任务", "状态", "进度", "计划起止"]):
        hdr[i].text = text
    for task in tasks:
        row = table.add_row().cells
        row[0].text = task.wbs_code
        row[1].text = task.name
        row[2].text = task.status
        row[3].text = _progress_text(task, effective)
        row[4].text = f"{task.plan_start or '-'} ~ {task.plan_end or '-'}"
    doc.add_heading("三、风险与问题", level=2)
    risks = list(db.scalars(select(Risk).where(Risk.project_id == project.id, Risk.archived_at.is_(None))))
    issues = list(db.scalars(select(Issue).where(Issue.project_id == project.id, Issue.status != "已关闭")))
    doc.add_paragraph(f"开放风险 {len(risks)} 项；未关闭问题 {len(issues)} 项")
    doc.add_paragraph("（详细台账请见系统）")
    return doc


def _build_final_doc(db: Session, project: Project) -> Document:
    doc = Document()
    doc.add_heading(f"项目收尾报告：{project.name}", level=1)
    doc.add_paragraph(f"项目编号：{project.code}    生成日期：{_utcnow().astimezone().date().isoformat()}")
    doc.add_paragraph(f"客户：{project.customer or '-'}    项目经理：{project.manager or '-'}")
    doc.add_paragraph(f"计划周期：{project.plan_start or '-'} ~ {project.plan_end or '-'}")

    tasks = wbs_service.list_tasks_ordered(db, project.id)
    metrics = progress_service.project_metrics(tasks)
    doc.add_heading("一、项目完成情况", level=2)
    doc.add_paragraph(
        f"任务完成：{metrics['completed_tasks']}/{metrics['total_tasks']}；整体进度 {metrics['progress']}%"
    )
    doc.add_paragraph(f"总工时 {metrics['total_hours']:.1f}h，剩余工时 {metrics['remaining_hours']:.1f}h")
    milestones = list(db.scalars(select(Milestone).where(Milestone.project_id == project.id)))
    doc.add_paragraph(f"里程碑 {len(milestones)} 项：{ '、'.join(m.name for m in milestones) or '-' }")

    doc.add_heading("二、交付资料清单", level=2)
    items = list(db.scalars(select(Deliverable).where(Deliverable.project_id == project.id)))
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    for i, text in enumerate(["类别", "资料名称", "提交状态", "审核状态"]):
        table.rows[0].cells[i].text = text
    for item in items:
        row = table.add_row().cells
        row[0].text = item.category or "-"
        row[1].text = item.name
        row[2].text = item.submit_status
        row[3].text = item.review_status

    doc.add_heading("三、风险与问题", level=2)
    risks = list(db.scalars(select(Risk).where(Risk.project_id == project.id, Risk.archived_at.is_(None))))
    issues = list(db.scalars(select(Issue).where(Issue.project_id == project.id)))
    doc.add_paragraph(f"风险 {len(risks)} 项、问题 {len(issues)} 项（明细见系统台账）")
    return doc


@router.get("/projects/{project_id}/stakeholders/report/export")
def export_stakeholder_report(project_id: UUID, db: Session = Depends(get_db)) -> StreamingResponse:
    """干系人管理报告 Word 导出（按单位→角色分类分组，含统计/矩阵/联系方式/参与程度）。"""
    from app.core.errors import AppError
    project = db.scalar(select(Project).where(Project.id == project_id))
    if project is None:
        raise AppError(404, "not_found", "项目不存在")
    stakeholders = list(db.scalars(select(Stakeholder).where(Stakeholder.project_id == project_id)))
    if not stakeholders:
        raise AppError(422, "empty", "该项目暂无干系人数据")

    doc = Document()
    doc.add_heading(f"干系人管理报告：{project.name}", level=1)
    now = _utcnow().astimezone()
    doc.add_paragraph(f"项目编号：{project.code}    导出日期：{now.date().isoformat()}")

    # 统计
    p0 = sum(1 for s in stakeholders if s.priority == "P0")
    p1 = sum(1 for s in stakeholders if s.priority == "P1")
    p2 = sum(1 for s in stakeholders if s.priority == "P2")
    hi_inf = sum(1 for s in stakeholders if s.influence >= 4)
    doc.add_heading("一、统计摘要", level=2)
    doc.add_paragraph(f"干系人总数：{len(stakeholders)}    P0 重点管理：{p0}    P1 重点关注：{p1}    P2 一般跟进：{p2}    高影响力(≥4)：{hi_inf}")

    # 按单位→角色分类分组
    org_map: dict[str, dict[str, list[Stakeholder]]] = {}
    for s in stakeholders:
        org = s.org or "未分类"
        sub = s.sub_pos or "未分类"
        if org not in org_map:
            org_map[org] = {}
        if sub not in org_map[org]:
            org_map[org][sub] = []
        org_map[org][sub].append(s)

    doc.add_heading("二、干系人明细（按单位分组）", level=2)
    for org, sub_map in org_map.items():
        total = sum(len(v) for v in sub_map.values())
        p = doc.add_paragraph()
        p.add_run(f"■ {org}（{total} 人）").bold = True
        for sub, people in sub_map.items():
            p2 = doc.add_paragraph()
            p2.paragraph_format.left_indent = 365760  # 0.4 inch indent (int for python-docx)
            p2.add_run(f"├ {sub}（{len(people)} 人）").bold = True
            table = doc.add_table(rows=1, cols=7)
            table.style = "Light Grid Accent 1"
            for i, text in enumerate(["姓名", "职位/角色", "优先级", "参与程度", "联系方式", "需求/诉求", "沟通计划"]):
                table.rows[0].cells[i].text = text
            for s in people:
                row = table.add_row().cells
                row[0].text = s.name
                row[1].text = s.role or "-"
                row[2].text = s.priority or "-"
                row[3].text = s.participation or "-"
                contact = []
                if s.contact_phone:
                    contact.append(f"📱{s.contact_phone}")
                if s.contact_wechat:
                    contact.append(f"💬{s.contact_wechat}")
                if s.contact_email:
                    contact.append(f"✉️{s.contact_email}")
                row[4].text = "\n".join(contact) or "-"
                row[5].text = s.demands or "-"
                row[6].text = s.comm_frequency or "-"
            doc.add_paragraph()  # 空行分隔

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    filename = f"stakeholders_report_{project.code}_{now.date().isoformat()}.docx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )