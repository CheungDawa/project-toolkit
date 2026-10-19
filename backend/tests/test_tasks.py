"""任务树 / WBS 编码 / 复制 / 删除 / 日期 / 工期 测试（A2-1/A2-9）。"""
import io
import uuid

from openpyxl import Workbook, load_workbook

from app.api.tasks import EXPORT_HEADERS

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _create_project(client, code="T-001") -> str:
    resp = client.post("/api/projects", json={"code": code, "name": "WBS 测试项目"})
    assert resp.status_code == 201
    return resp.json()["id"]


def _create_task(client, project_id, name, parent_id=None, **extra) -> dict:
    payload = {"name": name, "parent_id": parent_id, **extra}
    resp = client.post(f"/api/projects/{project_id}/tasks", json=payload)
    assert resp.status_code == 201, resp.text
    return resp.json()


def test_wbs_code_hierarchy(client):
    pid = _create_project(client)
    root_a = _create_task(client, pid, "A")
    child_b = _create_task(client, pid, "B", root_a["id"])
    child_c = _create_task(client, pid, "C", root_a["id"])
    root_d = _create_task(client, pid, "D")
    grand = _create_task(client, pid, "B-1", child_b["id"])

    listed = client.get(f"/api/projects/{pid}/tasks").json()["items"]
    by_name = {t["name"]: t for t in listed}
    assert by_name["A"]["wbs_code"] == "1"
    assert by_name["B"]["wbs_code"] == "1.1"
    assert by_name["C"]["wbs_code"] == "1.2"
    assert by_name["D"]["wbs_code"] == "2"
    assert by_name["B-1"]["wbs_code"] == "1.1.1"
    # 稳定 ID：uuid 不因编码重算变化
    assert by_name["A"]["id"] == root_a["id"]


def test_reorder_updates_wbs(client):
    pid = _create_project(client)
    a = _create_task(client, pid, "A")
    b = _create_task(client, pid, "B")
    c = _create_task(client, pid, "C")
    assert [t["wbs_code"] for t in client.get(f"/api/projects/{pid}/tasks").json()["items"]] == ["1", "2", "3"]

    resp = client.post(
        f"/api/projects/{pid}/tasks/reorder",
        json={"task_id": c["id"], "new_parent_id": None, "new_index": 0},
    )
    assert resp.status_code == 200
    codes = {t["name"]: t["wbs_code"] for t in client.get(f"/api/projects/{pid}/tasks").json()["items"]}
    assert codes == {"C": "1", "A": "2", "B": "3"}


def test_reorder_into_descendant_rejected(client):
    pid = _create_project(client)
    a = _create_task(client, pid, "A")
    child = _create_task(client, pid, "A-1", a["id"])
    resp = client.post(
        f"/api/projects/{pid}/tasks/reorder",
        json={"task_id": a["id"], "new_parent_id": child["id"], "new_index": 0},
    )
    assert resp.status_code == 422


def test_copy_subtree(client):
    pid = _create_project(client)
    a = _create_task(client, pid, "A")
    child = _create_task(client, pid, "A-1", a["id"])
    resp = client.post(f"/api/tasks/{a['id']}/copy")
    assert resp.status_code == 200
    listed = client.get(f"/api/projects/{pid}/tasks").json()["items"]
    names = [t["name"] for t in listed]
    assert "A（副本）" in names
    assert names.count("A-1") == 2


def test_delete_cascade(client):
    pid = _create_project(client)
    a = _create_task(client, pid, "A")
    _create_task(client, pid, "A-1", a["id"])
    _create_task(client, pid, "B")
    assert client.delete(f"/api/tasks/{a['id']}").status_code == 204
    listed = client.get(f"/api/projects/{pid}/tasks").json()["items"]
    assert [t["name"] for t in listed] == ["B"]
    assert listed[0]["wbs_code"] == "1"


def test_task_duration_workdays(client):
    pid = _create_project(client)
    task = _create_task(client, pid, "跨周末任务", plan_start="2026-10-01", plan_end="2026-10-05")
    # 2026-10-01(四)~10-05(一)：工作日 10-01,10-02,10-05 = 3 天
    assert task["duration_days"] == 3


def test_task_invalid_dates_rejected(client):
    pid = _create_project(client)
    resp = client.post(
        f"/api/projects/{pid}/tasks",
        json={"name": "坏日期", "plan_start": "2026-10-10", "plan_end": "2026-10-01"},
    )
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "date_invalid"


def test_parent_progress_rollup(client):
    pid = _create_project(client)
    parent = _create_task(client, pid, "父任务")
    _create_task(client, pid, "子1", parent["id"], progress=100, work_hours=2)
    _create_task(client, pid, "子2", parent["id"], progress=50, work_hours=2)
    listed = client.get(f"/api/projects/{pid}/tasks").json()["items"]
    parent_read = next(t for t in listed if t["name"] == "父任务")
    assert parent_read["effective_progress"] == 75.0


def test_parent_progress_equal_fallback(client):
    pid = _create_project(client)
    parent = _create_task(client, pid, "父任务")
    _create_task(client, pid, "子1", parent["id"], progress=100)
    _create_task(client, pid, "子2", parent["id"], progress=0)
    listed = client.get(f"/api/projects/{pid}/tasks").json()["items"]
    parent_read = next(t for t in listed if t["name"] == "父任务")
    assert parent_read["effective_progress"] == 50.0


def test_leaf_not_double_counted(client):
    """父子混合统计只计叶子（甘特条8）：父任务自身进度不参与自身汇总。"""
    pid = _create_project(client)
    parent = _create_task(client, pid, "父任务", progress=100)  # 父任务自身进度应被忽略
    _create_task(client, pid, "子1", parent["id"], progress=100)
    listed = client.get(f"/api/projects/{pid}/tasks").json()["items"]
    parent_read = next(t for t in listed if t["name"] == "父任务")
    assert parent_read["effective_progress"] == 100.0  # 来自子任务而非自身
    leaf = next(t for t in listed if t["name"] == "子1")
    assert leaf["effective_progress"] == 100.0


def test_tasks_persist_after_reload(client):
    """刷新等价重查（A2-9）。"""
    pid = _create_project(client)
    _create_task(client, pid, "A")
    first = client.get(f"/api/projects/{pid}/tasks").json()
    second = client.get(f"/api/projects/{pid}/tasks").json()
    assert first == second


# ---------- Excel 导入（P0-2 重写） ----------

def _upload_import(client, project_id, data_rows, header=None):
    """构造与导出同构的 xlsx 并 POST 导入；返回响应对象。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(header or EXPORT_HEADERS))
    for row in data_rows:
        sheet.append(list(row))
    buf = io.BytesIO()
    workbook.save(buf)
    return client.post(
        f"/api/projects/{project_id}/tasks/import",
        files={"file": ("tasks.xlsx", buf.getvalue(), XLSX_MIME)},
    )


def test_import_multilevel_tree(client):
    """多层文件一次导入：树结构、编码、父子关系、工期全部正确（原实现 100% 失败的场景）。"""
    pid = _create_project(client)
    resp = _upload_import(client, pid, [
        ("1", "需求调研", "张三", "进行中", "高", "2026-10-01", "2026-10-05", 40, 30, "调研报告"),
        ("1.1", "现场访谈", "李四", "未开始", "中", "", "", 16, 0, ""),
        ("1.1.1", "访谈提纲", "李四", "未开始", "中", "", "", 4, 0, ""),
        ("2", "方案设计", "王五", "未开始", "高", "", "", 24, 0, ""),
    ])
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["created"] == 4
    assert body["errors"] == []
    assert body["warnings"] == []

    items = client.get(f"/api/projects/{pid}/tasks").json()["items"]
    by_name = {t["name"]: t for t in items}
    assert {name: t["wbs_code"] for name, t in by_name.items()} == {
        "需求调研": "1", "现场访谈": "1.1", "访谈提纲": "1.1.1", "方案设计": "2",
    }
    assert by_name["现场访谈"]["parent_id"] == by_name["需求调研"]["id"]
    assert by_name["访谈提纲"]["parent_id"] == by_name["现场访谈"]["id"]
    assert by_name["方案设计"]["parent_id"] is None
    assert by_name["需求调研"]["sort_order"] == 0
    assert by_name["现场访谈"]["sort_order"] == 0
    assert by_name["方案设计"]["sort_order"] == 1
    # 工期按工作日历计算（与创建任务 API 同语义：2026-10-01~10-05 = 3 个工作日）
    assert by_name["需求调研"]["duration_days"] == 3


def test_import_export_roundtrip(client):
    """导出 → 导入到新项目 → 再导出：(wbs_code, name) 完全一致（树结构不变）。"""
    pid = _create_project(client)
    root = _create_task(client, pid, "开发")
    _create_task(client, pid, "前端", root["id"])
    _create_task(client, pid, "后端", root["id"])
    root2 = _create_task(client, pid, "测试")
    case = _create_task(client, pid, "用例", root2["id"])
    _create_task(client, pid, "冒烟", case["id"])

    first_export = client.get(f"/api/projects/{pid}/tasks/export")
    assert first_export.status_code == 200

    pid2 = _create_project(client, "T-IMP-2")
    imported = client.post(
        f"/api/projects/{pid2}/tasks/import",
        files={"file": ("tasks.xlsx", first_export.content, XLSX_MIME)},
    )
    assert imported.status_code == 200, imported.text
    assert imported.json()["created"] == 6
    assert imported.json()["errors"] == []

    second_export = client.get(f"/api/projects/{pid2}/tasks/export")
    assert second_export.status_code == 200

    def _rows(content):
        workbook = load_workbook(io.BytesIO(content))
        return [
            (row[0], row[1])
            for row in workbook.active.iter_rows(min_row=2, values_only=True)
        ]

    assert _rows(first_export.content) == _rows(second_export.content)


def test_import_invalid_progress_is_row_error(client):
    """progress=150 的非法行返回行级错误而不是 5xx，且不拖累同批其他行。"""
    pid = _create_project(client)
    resp = _upload_import(client, pid, [
        ("1", "正常任务", "", "", "", "", "", "", 0, ""),
        ("2", "进度越界", "", "", "", "", "", "", 150, ""),
    ])
    assert resp.status_code == 200, resp.text  # 不再是 IntegrityError 引发的 500
    body = resp.json()
    assert body["created"] == 1
    assert len(body["errors"]) == 1
    assert body["errors"][0]["row"] == 3  # 表头第 1 行，第 2 条数据在第 3 行
    assert "0 到 100" in body["errors"][0]["message"]
    assert "150" in body["errors"][0]["message"]
    items = client.get(f"/api/projects/{pid}/tasks").json()["items"]
    assert [t["name"] for t in items] == ["正常任务"]


def test_import_template_download(client):
    """模板端点：200 + xlsx 类型 + 表头与导入解析逐字一致。"""
    pid = _create_project(client)
    resp = client.get(f"/api/projects/{pid}/tasks/import-template")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(XLSX_MIME)
    disposition = resp.headers.get("content-disposition", "")
    assert "filename*=UTF-8''" in disposition and disposition.endswith(".xlsx")

    workbook = load_workbook(io.BytesIO(resp.content))
    rows = list(workbook.active.iter_rows(values_only=True))
    assert list(rows[0]) == list(EXPORT_HEADERS)
    assert len(rows) >= 3  # 表头 + 示例数据行


# ---------- 跨项目克隆 WBS ----------

def test_clone_wbs_from_project(client):
    """整树深拷贝到新项目：层级保持、追加在已有任务之后、业务状态重置、不带日期。"""
    source_pid = _create_project(client, "T-CLN-1")
    root = _create_task(
        client, source_pid, "开发",
        plan_start="2026-10-01", plan_end="2026-10-05",
        work_hours=10, progress=80, status="已完成",
    )
    _create_task(client, source_pid, "前端", root["id"], work_hours=4)
    _create_task(client, source_pid, "后端", root["id"], work_hours=6, priority="高")

    target_pid = _create_project(client, "T-CLN-2")
    _create_task(client, target_pid, "既有任务")  # 目标已有任务不被覆盖

    resp = client.post(f"/api/projects/{target_pid}/tasks/clone-from/{source_pid}")
    assert resp.status_code == 200, resp.text
    assert resp.json() == {"created": 3, "with_dates": False}

    items = client.get(f"/api/projects/{target_pid}/tasks").json()["items"]
    by_name = {t["name"]: t for t in items}
    # 克隆树作为新的顶级分支接在既有任务之后，编号连续
    assert by_name["既有任务"]["wbs_code"] == "1"
    assert by_name["开发"]["wbs_code"] == "2"
    assert by_name["前端"]["wbs_code"] == "2.1"
    assert by_name["后端"]["wbs_code"] == "2.2"
    assert by_name["前端"]["parent_id"] == by_name["开发"]["id"]
    assert by_name["后端"]["parent_id"] == by_name["开发"]["id"]
    # 模板字段保留
    assert by_name["后端"]["priority"] == "高"
    assert by_name["开发"]["work_hours"] == 10
    # 业务状态重置 + 默认不带日期
    assert by_name["开发"]["status"] == "未开始"
    assert by_name["开发"]["progress"] == 0
    assert by_name["开发"]["plan_start"] is None
    assert by_name["开发"]["duration_days"] is None
    # 源项目不受影响
    source_items = client.get(f"/api/projects/{source_pid}/tasks").json()["items"]
    assert len(source_items) == 3


def test_clone_wbs_with_dates(client):
    """with_dates=true：连计划日期一起复制，并按工作日历重算工期。"""
    source_pid = _create_project(client, "T-CLN-3")
    _create_task(client, source_pid, "调研", plan_start="2026-10-01", plan_end="2026-10-05")
    target_pid = _create_project(client, "T-CLN-4")

    resp = client.post(f"/api/projects/{target_pid}/tasks/clone-from/{source_pid}?with_dates=true")
    assert resp.status_code == 200, resp.text
    assert resp.json() == {"created": 1, "with_dates": True}

    items = client.get(f"/api/projects/{target_pid}/tasks").json()["items"]
    assert items[0]["plan_start"] == "2026-10-01"
    assert items[0]["plan_end"] == "2026-10-05"
    assert items[0]["duration_days"] == 3  # 与创建任务 API 同口径（工作日历）


def test_clone_wbs_source_not_found(client):
    """源项目不存在 → 404 not_found。"""
    pid = _create_project(client, "T-CLN-5")
    resp = client.post(f"/api/projects/{pid}/tasks/clone-from/{uuid.uuid4()}")
    assert resp.status_code == 404
    assert resp.json()["error"]["code"] == "not_found"